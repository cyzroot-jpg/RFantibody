"""Operator design document schema v3 for teco-ops operator delivery."""
from __future__ import annotations

import json
import re
import warnings
from copy import deepcopy
from dataclasses import dataclass
from datetime import date
from functools import lru_cache
from pathlib import Path
from typing import Any

SCHEMA_VERSION = "3.0"
OP_RE = re.compile(r"^[a-z][a-z0-9_]*$")
DESIGN_DIR_NAME = "op_design"

DEFAULT_LDM_BUDGET_BYTES = 256 * 1024
DEFAULT_LDM_USABLE_BYTES = 235 * 1024

VALID_COMPUTATION_TYPES = frozenset(
    {"elementwise", "reduction", "normalization", "matmul", "custom"}
)
VALID_BOTTLENECK_TYPES = frozenset(
    {"compute", "memory", "latency", "mixed", "unknown"}
)
STRICT_BOTTLENECK_TYPES = frozenset({"compute", "memory", "latency", "mixed"})
VALID_SECTION_STATUS = frozenset({"draft", "defined"})
VALID_SCAFFOLD_FAMILIES = frozenset({"elewise", "reduce", "norm", "manual"})
VALID_WORKFLOW_FOCUS = frozenset(
    {
        "operator_definition",
        "throughput_model",
        "tiling",
        "core_partition",
        "double_buffer",
        "ldm",
        "pipeline",
        "complete",
    }
)

VALID_DOC_STATUS = frozenset({"draft", "iterating", "review", "approved"})

STANDARD_TOP_LEVEL_KEYS = frozenset(
    {
        "schema_version",
        "op",
        "title",
        "status",
        "created_at",
        "updated_at",
        "workflow",
        "operator_definition",
        "flops",
        "parameters",
        "throughput_model",
        "tiling_strategy",
        "core_partition_strategy",
        "double_buffer_strategy",
        "ldm_planning",
        "pipeline_layout",
        "scaffold",
        "golden",
    }
)

SECTIONS_REQUIRING_STATUS = (
    "tiling_strategy",
    "core_partition_strategy",
    "double_buffer_strategy",
    "pipeline_layout",
)

COMPUTATION_TYPE_TO_SCAFFOLD: dict[str, str] = {
    "elementwise": "elewise",
    "reduction": "reduce",
    "normalization": "norm",
    "matmul": "manual",
    "custom": "manual",
}

WORKFLOW_ORDER: list[tuple[str, str, bool]] = [
    ("operator_definition", "1. 算子定义", True),
    ("flops", "1. 计算 FLOPS（与算子定义同优先级）", True),
    ("parameters", "2. 参数", True),
    ("throughput_model", "2. 计算吞吐 / 瓶颈（与参数同优先级）", True),
    ("tiling_strategy", "3. 分块策略", False),
    ("core_partition_strategy", "4. 分核策略", False),
    ("double_buffer_strategy", "5. 双缓冲策略", False),
    ("ldm_planning", "6. LDM 使用规划", True),
    ("pipeline_layout", "7. 流水排布", False),
]

T1_MATMUL_CAPS = {
    "cluster_ace_peak_tflops": 75.0,
    "cluster_dma_read_gbps": 313.0,
    "cluster_dma_write_gbps": 270.0,
    "cluster_rma_p2p_gbps": 53.0,
    "cluster_rma_row_bcast_gbps": 188.0,
    "cluster_rma_col_bcast_gbps": 278.0,
    "ldm_bank_gbps": 140.0,
    "ace_west_gbps": 80.0,
    "ace_north_gbps": 80.0,
    "ace_writeback_gbps": 27.0,
    "spe_clock_hz": 2500000000,
}

REFERENCE_DIR = Path(__file__).resolve().parents[1] / "references"
SOURCE_DOCS_DIR = REFERENCE_DIR / "source-docs"
ACE_XLSX_PATH = SOURCE_DOCS_DIR / "ACE.xlsx"
DMA_XLSX_PATH = SOURCE_DOCS_DIR / "DMA测试.xlsx"


@dataclass(frozen=True)
class DesignPaths:
    demo_root: Path
    op: str

    @property
    def op_dir(self) -> Path:
        return self.demo_root / DESIGN_DIR_NAME / self.op

    @property
    def json_path(self) -> Path:
        return self.op_dir / "design.json"

    @property
    def md_path(self) -> Path:
        return self.op_dir / "design.md"


def design_paths(demo_root: Path, op: str) -> DesignPaths:
    return DesignPaths(demo_root=demo_root.resolve(), op=op)


def validate_op_name(op: str) -> None:
    if not OP_RE.match(op):
        raise ValueError(f"invalid op '{op}': use snake_case [a-z][a-z0-9_]*")


def _default_shapes(computation_type: str, *, smoke: bool) -> dict[str, int]:
    if computation_type == "elementwise":
        return {"N": 8192 if smoke else 4096, "H": 2, "W": 2, "C": 1024}
    if computation_type == "matmul":
        return (
            {"M": 64, "N": 32, "K": 128}
            if smoke
            else {"M": 512, "N": 256, "K": 8192}
        )
    if computation_type == "normalization":
        return (
            {"N": 4, "H": 4, "W": 8, "C": 8192}
            if smoke
            else {"N": 8, "H": 8, "W": 16, "C": 8192}
        )
    return {"N": 2, "H": 2, "W": 2, "C": 1024}


def _default_operator_family(computation_type: str) -> str:
    if computation_type == "matmul":
        return "gemm"
    if computation_type == "normalization":
        return "rowwise_norm"
    if computation_type == "reduction":
        return "rowwise_reduce"
    if computation_type == "elementwise":
        return "elementwise"
    return "custom"


def default_design(
    op: str,
    *,
    computation_type: str = "elementwise",
    title: str = "",
) -> dict[str, Any]:
    validate_op_name(op)
    if computation_type not in VALID_COMPUTATION_TYPES:
        raise ValueError(f"unknown computation_type: {computation_type}")

    smoke = _default_shapes(computation_type, smoke=True)
    default = _default_shapes(computation_type, smoke=False)
    scaffold_family = COMPUTATION_TYPE_TO_SCAFFOLD.get(computation_type, "elewise")
    dtype_io = "float16" if computation_type in {"normalization", "matmul"} else "float32"
    operator_family = _default_operator_family(computation_type)

    return {
        "schema_version": SCHEMA_VERSION,
        "op": op,
        "title": title or op.replace("_", " ").title(),
        "status": "draft",
        "created_at": date.today().isoformat(),
        "updated_at": date.today().isoformat(),
        "workflow": {
            "mode": "iterative",
            "description": "非单链路；循环往复，不断接近理论峰值",
            "current_focus": "operator_definition",
            "iteration_round": 1,
            "history": [],
        },
        "operator_definition": {
            "priority": "P0",
            "status": "draft",
            "summary": "",
            "computation_type": computation_type,
            "operator_family": operator_family,
            "formula": "",
            "logic_steps": [],
            "io_tensors": [],
            "kernel_apis": [],
            "independent_axes": [],
            "reduction_axes": [],
            "reuse_axes": [],
        },
        "flops": {
            "priority": "P0",
            "status": "draft",
            "formula_breakdown": [],
            "flops_per_element": None,
            "total_flops": None,
            "notes": "按公式逐项统计乘加次数",
        },
        "parameters": {
            "priority": "P0",
            "status": "draft",
            "shape_mode": "locked",
            "shape_keys": list(smoke.keys()),
            "shapes": {"smoke": smoke, "default": default},
            "locked_shapes": deepcopy(default),
            "shape_ranges": {},
            "ideal_shape_goal": {
                "enabled": False,
                "optimize_for": "throughput",
                "fixed_axes": {},
                "search_axes": [],
                "notes": "",
            },
            "dtype_io": dtype_io,
            "dtype_acc": "float32",
            "hyperparams": {},
        },
        "throughput_model": {
            "priority": "P0",
            "status": "draft",
            "model_type": "layered_t1",
            "spe_clock_hz": None,
            "peak_flops_per_spe": None,
            "hbm_bandwidth_gbps": None,
            "arithmetic_intensity": None,
            "theoretical_peak_tflops": None,
            "upper_bound_tflops": None,
            "hardware_caps": {
                "ace_peak_tflops": None,
                "dma_read_gbps": None,
                "dma_write_gbps": None,
                "rma_p2p_gbps": None,
                "rma_row_bcast_gbps": None,
                "rma_col_bcast_gbps": None,
                "ldm_bank_gbps": None,
                "ace_west_gbps": None,
                "ace_north_gbps": None,
                "ace_writeback_gbps": None,
            },
            "traffic_model": {
                "hbm_read_bytes": None,
                "hbm_write_bytes": None,
                "rma_bytes": None,
                "compute_flops": None,
            },
            "stage_time_model_us": {
                "dma": None,
                "rma": None,
                "ace": None,
                "simd": None,
                "compute_group": None,
                "writeback": None,
                "sync": None,
            },
            "constraint_discounts": [],
            "achievable_throughput_note": "",
            "bottleneck_type": "unknown",
            "bottleneck_analysis": "",
            "target_efficiency": None,
        },
        "tiling_strategy": {
            "status": "draft",
            "tile_axis": "cols",
            "tile_elems": 4096,
            "dma_bytes_per_tile": 8192,
            "alignment_bytes": 4,
            "hierarchy": [],
            "selected_tile": {},
            "search_space": {},
            "constraints": [],
            "feasible_region_summary": "",
            "rationale": "",
        },
        "core_partition_strategy": {
            "status": "draft",
            "slave_core_num": 32,
            "partition_axis": "rows",
            "strategy": "stride",
            "pattern": "for (row = _PEN; row < rows; row += slave_core_num)",
            "grid_shape": {},
            "split_dimensions": [],
            "communication_pattern": [],
            "reduction_pattern": "",
            "loader_assignment": [],
            "rationale": "",
        },
        "double_buffer_strategy": {
            "status": "draft",
            "enabled": False,
            "buffered_tensors": [],
            "buffer_depth": 2,
            "overlap_dma_compute": False,
            "stage_buffers": [],
            "enable_condition": "",
            "blocked_by": [],
            "rationale": "",
        },
        "ldm_planning": {
            "priority": "P0",
            "status": "draft",
            "budget_bytes_per_spe": DEFAULT_LDM_BUDGET_BYTES,
            "usable_budget_bytes": DEFAULT_LDM_USABLE_BYTES,
            "buffers": [],
            "total_bytes": 0,
            "fits_budget": True,
            "bank_allocation": [],
            "pressure_summary": "",
            "reuse_notes": "",
            "allocation_notes": "",
        },
        "pipeline_layout": {
            "status": "draft",
            "passes": [],
            "dataflow_description": "",
            "sync_points": [],
            "steady_state_stages": [],
            "overlap_pairs": [],
            "wait_points": [],
            "fill_drain_notes": "",
        },
        "scaffold": {
            "family": scaffold_family,
            "smoke": True,
            "patch_runsh": True,
            "upgrade_common": True,
        },
        "golden": {
            "framework": "pytorch",
            "atol": 3e-3 if computation_type == "normalization" else 1e-5,
            "rtol": 2e-2 if computation_type == "normalization" else 3e-2,
        },
    }


def _ceil_div(a: int, b: int) -> int:
    return (a + b - 1) // b


def _parse_int_label(value: Any) -> int | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value)
    match = re.search(r"-?\d+", str(value))
    return int(match.group(0)) if match else None


def _to_float(value: Any) -> float | None:
    if value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _nearest_key(keys: list[int], value: int) -> int | None:
    if not keys:
        return None
    return min(keys, key=lambda key: abs(key - value))


def _interp_1d(points: dict[int, float], x: int) -> float | None:
    if not points:
        return None
    keys = sorted(points)
    if x in points:
        return points[x]
    if x <= keys[0]:
        return points[keys[0]]
    if x >= keys[-1]:
        return points[keys[-1]]
    lo = keys[0]
    hi = keys[-1]
    for idx in range(len(keys) - 1):
        if keys[idx] <= x <= keys[idx + 1]:
            lo, hi = keys[idx], keys[idx + 1]
            break
    if lo == hi:
        return points[lo]
    return points[lo] + (points[hi] - points[lo]) * (x - lo) / (hi - lo)


def _interp_2d(table: dict[int, dict[int, float]], x: int, y: int) -> float | None:
    if not table:
        return None
    xs = sorted(table)
    if x <= xs[0]:
        x0 = x1 = xs[0]
    elif x >= xs[-1]:
        x0 = x1 = xs[-1]
    else:
        x0, x1 = xs[0], xs[-1]
        for idx in range(len(xs) - 1):
            if xs[idx] <= x <= xs[idx + 1]:
                x0, x1 = xs[idx], xs[idx + 1]
                break
    v0 = _interp_1d(table.get(x0, {}), y)
    v1 = _interp_1d(table.get(x1, {}), y)
    if v0 is None:
        return v1
    if v1 is None:
        return v0
    if x0 == x1:
        return v0
    return v0 + (v1 - v0) * (x - x0) / (x1 - x0)


def _load_matrix_sheet(workbook: Any, sheet_name: str) -> dict[int, dict[int, float]]:
    ws = workbook[sheet_name]
    k_cols: dict[int, int] = {}
    for col in range(2, ws.max_column + 1):
        k_val = _parse_int_label(ws.cell(1, col).value)
        if k_val is not None:
            k_cols[col] = k_val
    table: dict[int, dict[int, float]] = {}
    for row in range(2, ws.max_row + 1):
        m_val = _parse_int_label(ws.cell(row, 1).value)
        if m_val is None:
            continue
        row_values: dict[int, float] = {}
        for col, k_val in k_cols.items():
            v = _to_float(ws.cell(row, col).value)
            if v is not None:
                row_values[k_val] = v
        if row_values:
            table[m_val] = row_values
    return table


def _load_dma_contiguous_sheet(workbook: Any) -> dict[int, float]:
    ws = workbook["dma 连续"]
    points: dict[int, float] = {}
    for row in range(5, ws.max_row + 1):
        length_kb = _parse_int_label(ws.cell(row, 4).value)
        bw = _to_float(ws.cell(row, 6).value)
        if length_kb is not None and bw is not None:
            points[length_kb * 1024] = bw
    return points


def _load_dma_stride_sheet(workbook: Any, sheet_name: str) -> list[dict[str, float]]:
    ws = workbook[sheet_name]
    rows: list[dict[str, float]] = []
    for start_col in range(2, ws.max_column + 1, 5):
        for row in range(5, ws.max_row + 1):
            bm = _parse_int_label(ws.cell(row, start_col).value)
            bk = _parse_int_label(ws.cell(row, start_col + 1).value)
            lda = _parse_int_label(ws.cell(row, start_col + 2).value)
            bw = _to_float(ws.cell(row, start_col + 3).value)
            if bm is None or bk is None or lda is None or bw is None:
                continue
            rows.append({"bm": float(bm), "bk": float(bk), "lda": float(lda), "bw": float(bw)})
    return rows


@lru_cache(maxsize=1)
def _load_calibration_tables() -> dict[str, Any]:
    tables: dict[str, Any] = {
        "available": False,
        "notes": [],
        "ace_tflops": {},
        "ace_half_tflops": {},
        "ace_delay_cycles": {},
        "dma_contiguous": {},
        "dma_stride": {},
    }
    try:
        from openpyxl import load_workbook
    except ImportError:
        tables["notes"].append("openpyxl unavailable; using built-in fallback model")
        return tables

    if ACE_XLSX_PATH.is_file():
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                ace_wb = load_workbook(ACE_XLSX_PATH, data_only=True, read_only=False)
            tables["ace_tflops"] = _load_matrix_sheet(ace_wb, "TFLOPS")
            tables["ace_half_tflops"] = _load_matrix_sheet(ace_wb, "write half")
            tables["ace_delay_cycles"] = _load_matrix_sheet(ace_wb, "下发_delay")
            tables["notes"].append(f"ACE calibration loaded from {ACE_XLSX_PATH.name}")
        except Exception as exc:  # pragma: no cover - depends on workbook parser
            tables["notes"].append(f"ACE calibration load failed: {exc}")
    else:
        tables["notes"].append(f"ACE calibration missing: {ACE_XLSX_PATH}")

    if DMA_XLSX_PATH.is_file():
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                dma_wb = load_workbook(DMA_XLSX_PATH, data_only=True, read_only=False)
            tables["dma_contiguous"] = _load_dma_contiguous_sheet(dma_wb)
            stride = _load_dma_stride_sheet(dma_wb, "2.5Ghz - dma 跨步")
            if not stride:
                stride = _load_dma_stride_sheet(dma_wb, "2.36Ghz - dma 跨步")
            tables["dma_stride"] = stride
            tables["notes"].append(f"DMA calibration loaded from {DMA_XLSX_PATH.name}")
        except Exception as exc:  # pragma: no cover - depends on workbook parser
            tables["notes"].append(f"DMA calibration load failed: {exc}")
    else:
        tables["notes"].append(f"DMA calibration missing: {DMA_XLSX_PATH}")

    tables["available"] = bool(tables["ace_tflops"] or tables["dma_contiguous"] or tables["dma_stride"])
    return tables


def _ace_calibrated_tflops(m2: int, k2: int, *, output_dtype: str = "float32") -> float | None:
    tables = _load_calibration_tables()
    table_key = "ace_half_tflops" if output_dtype in {"float16", "half", "fp16"} else "ace_tflops"
    value = _interp_2d(tables.get(table_key, {}), max(1, min(128, m2)), max(32, k2))
    return float(value) if value is not None else None


def _ace_delay_us(m2: int, k2: int, caps: dict[str, Any]) -> float:
    tables = _load_calibration_tables()
    cycles = _interp_2d(tables.get("ace_delay_cycles", {}), max(1, min(128, m2)), max(32, k2))
    if cycles is None:
        return 0.0
    ace_hz = float(caps.get("spe_clock_hz", 2500000000)) / 2.0
    return float(cycles) / ace_hz * 1e6


def _dma_stride_gbps_from_hint(hint: dict[str, int]) -> float | None:
    rows = _load_calibration_tables().get("dma_stride", [])
    if not rows:
        return None
    bm = int(hint.get("bm", 0))
    bk = int(hint.get("bk", 0))
    lda = int(hint.get("lda", 0))
    if bm <= 0 or bk <= 0 or lda <= 0:
        return None

    same_shape = [row for row in rows if int(row["bm"]) == bm and int(row["bk"]) == bk]
    if same_shape:
        by_lda = {int(row["lda"]): float(row["bw"]) for row in same_shape}
        return _interp_1d(by_lda, lda)

    def score(row: dict[str, float]) -> float:
        return (
            abs(row["bm"] - bm) / max(1, bm)
            + abs(row["bk"] - bk) / max(1, bk)
            + abs(row["lda"] - lda) / max(1, lda)
        )

    return float(min(rows, key=score)["bw"])


def _dma_calibrated_gbps(
    bytes_per_transfer: int,
    *,
    stride: bool = False,
    hint: dict[str, int] | None = None,
) -> float | None:
    tables = _load_calibration_tables()
    if stride:
        if hint:
            hinted_bw = _dma_stride_gbps_from_hint(hint)
            if hinted_bw is not None:
                return hinted_bw
        # Fallback for callers without enough dataflow information: nearest
        # observed transfer shape, not the optimistic best for the same size.
        rows = tables.get("dma_stride", [])
        if rows:
            target = max(1, bytes_per_transfer)
            return float(min(rows, key=lambda row: abs(row["bm"] * row["bk"] * 2 - target))["bw"])
    bw = _interp_1d(tables.get("dma_contiguous", {}), max(1, bytes_per_transfer))
    return float(bw) if bw is not None else None


def _dma_read_time_us(
    bytes_total: int,
    bytes_per_transfer: int,
    caps: dict[str, Any],
    *,
    stride: bool = False,
    hint: dict[str, int] | None = None,
) -> float:
    bw = _dma_calibrated_gbps(bytes_per_transfer, stride=stride, hint=hint)
    if bw is None or bw <= 0:
        bw = float(caps["cluster_dma_read_gbps"])
    return bytes_total / (bw * 1e9) * 1e6


def _dma_read_segments_time_us(
    segments: list[dict[str, Any]],
    caps: dict[str, Any],
) -> float:
    total = 0.0
    for segment in segments:
        total += _dma_read_time_us(
            int(segment["bytes_total"]),
            int(segment["bytes_per_transfer"]),
            caps,
            stride=bool(segment.get("stride", False)),
            hint=segment.get("hint"),
        )
    return total


def _ace_time_us(flops: int, m2: int, k2: int, caps: dict[str, Any], *, eff_scale: float = 1.0) -> float:
    calibrated = _ace_calibrated_tflops(m2, k2)
    if calibrated is not None and calibrated > 0:
        return flops / (calibrated * eff_scale * 1e12) * 1e6
    ace_eff = _estimate_ace_efficiency(m2, k2) * eff_scale
    return flops / (caps["cluster_ace_peak_tflops"] * max(0.08, ace_eff) * 1e12) * 1e6


SIMD_VECTOR_LANES = 16
SIMD_UNROLL_VECTORS = 8
SIMD_P0_LATENCY_CYCLES = {
    "fp32_add": 5.0,
    "fp32_sub": 5.0,
    "fp32_mul": 5.0,
    "fp32_fma": 5.0,
    "fp32_compare": 5.0,
    "fp32_select": 1.0,
    "int_add": 3.0,
    "int_mul": 3.0,
    "convert": 2.0,
    "permute": 3.0,
    "logic": 1.0,
}
SIMD_P1_LATENCY_CYCLES = {
    "load": 5.0,
    "store": 4.0,
    "branch": 1.0,
    "sync": 1.0,
}


def _simd_kernel_cycles(elements: int, kernel: dict[str, float]) -> dict[str, float]:
    if elements <= 0:
        return {
            "vectors": 0.0,
            "p0_issue_cycles": 0.0,
            "p1_issue_cycles": 0.0,
            "latency_cycles": 0.0,
            "tail_cycles": 0.0,
            "total_cycles": 0.0,
        }

    vector_lanes = int(kernel.get("vector_lanes", SIMD_VECTOR_LANES))
    vectors = _ceil_div(elements, max(1, vector_lanes))
    loads = float(kernel.get("loads", 0.0))
    stores = float(kernel.get("stores", 0.0))
    branches = float(kernel.get("branches", 0.0))
    syncs = float(kernel.get("syncs", 0.0))
    p0_ops = {
        key: float(kernel.get(key, 0.0))
        for key in SIMD_P0_LATENCY_CYCLES
    }

    p0_issue_per_vector = sum(p0_ops.values())
    p1_issue_per_vector = loads + stores + branches + syncs
    p0_issue_cycles = vectors * p0_issue_per_vector
    p1_issue_cycles = vectors * p1_issue_per_vector

    active_p0_latencies = [
        SIMD_P0_LATENCY_CYCLES[key]
        for key, count in p0_ops.items()
        if count > 0
    ]
    active_p1_latencies = []
    if loads > 0:
        active_p1_latencies.append(SIMD_P1_LATENCY_CYCLES["load"])
    if stores > 0:
        active_p1_latencies.append(SIMD_P1_LATENCY_CYCLES["store"])
    if branches > 0:
        active_p1_latencies.append(SIMD_P1_LATENCY_CYCLES["branch"])
    if syncs > 0:
        active_p1_latencies.append(SIMD_P1_LATENCY_CYCLES["sync"])

    latency_cycles = max(active_p0_latencies + active_p1_latencies, default=0.0)
    if loads > 0 and p0_issue_per_vector > 0:
        latency_cycles = max(latency_cycles, SIMD_P1_LATENCY_CYCLES["load"])
    if stores > 0 and p0_issue_per_vector > 0:
        latency_cycles = max(latency_cycles, max(active_p0_latencies, default=0.0))

    unroll_vectors = max(1, int(kernel.get("unroll_vectors", SIMD_UNROLL_VECTORS)))
    blocks = _ceil_div(vectors, unroll_vectors)
    tail_cycles = blocks * float(kernel.get("tail_cycles_per_unroll", 2.0))
    control_cycles = blocks * float(kernel.get("control_cycles_per_unroll", 1.0))
    total_cycles = max(p0_issue_cycles, p1_issue_cycles) + latency_cycles + tail_cycles + control_cycles
    return {
        "vectors": float(vectors),
        "p0_issue_cycles": p0_issue_cycles,
        "p1_issue_cycles": p1_issue_cycles,
        "latency_cycles": latency_cycles,
        "tail_cycles": tail_cycles,
        "control_cycles": control_cycles,
        "total_cycles": total_cycles,
    }


def _simd_time_us(
    *,
    elements: int,
    kernel: dict[str, float],
    caps: dict[str, Any] | None = None,
) -> float:
    if elements <= 0:
        return 0.0
    cycles = _simd_kernel_cycles(elements, kernel)["total_cycles"]
    hz = float((caps or T1_MATMUL_CAPS).get("spe_clock_hz", 2500000000))
    return cycles / hz * 1e6


def _simd_epilogue_time_us(elements: int, caps: dict[str, Any]) -> float:
    # Current GEMM output is fp32 C, so the default epilogue is only local
    # bookkeeping/tail handling. Bias/activation/cast paths can raise this.
    return _simd_time_us(
        elements=elements,
        kernel={
            "fp32_select": 0.10,
            "branches": 0.02,
            "control_cycles_per_unroll": 1.0,
            "tail_cycles_per_unroll": 1.0,
        },
        caps=caps,
    )


def _simd_workspace_reduce_time_us(elements: int, caps: dict[str, Any]) -> float:
    # Accumulate partial fp32 C from workspace: load acc + load partial +
    # vector add + store. P0 add and P1 load/store overlap through software
    # pipelining; the exposed cost is max(P0 issue, P1 issue) plus fill/drain.
    return _simd_time_us(
        elements=elements,
        kernel={
            "loads": 2.0,
            "stores": 1.0,
            "fp32_add": 1.0,
            "control_cycles_per_unroll": 1.0,
            "tail_cycles_per_unroll": 2.0,
        },
        caps=caps,
    )


def _exposed_stage_times(t_dma: float, t_rma: float, t_ace: float, t_simd: float, t_wb: float) -> list[tuple[str, float]]:
    return [
        ("dma", t_dma),
        ("rma", t_rma),
        ("compute_group", max(t_ace, t_simd)),
        ("writeback", t_wb),
    ]


def _candidate_bottleneck_stage(t_dma: float, t_rma: float, t_ace: float, t_simd: float, t_wb: float) -> str:
    stage, _ = max(_exposed_stage_times(t_dma, t_rma, t_ace, t_simd, t_wb), key=lambda item: item[1])
    if stage != "compute_group":
        return stage
    return "simd" if t_simd > t_ace else "ace"


def _ace_m_efficiency(m2: int) -> float:
    if m2 <= 16:
        return 0.17
    if m2 <= 32:
        return 0.34
    if m2 <= 64:
        return 0.67
    if m2 <= 96:
        return 0.82
    return 1.0


def _ace_k_efficiency(k2: int) -> float:
    if k2 <= 32:
        return 0.16
    if k2 <= 64:
        return 0.28
    if k2 <= 96:
        return 0.40
    if k2 <= 128:
        return 0.52
    if k2 <= 192:
        return 0.68
    if k2 <= 256:
        return 0.78
    if k2 <= 384:
        return 0.89
    return 1.0


def _estimate_ace_efficiency(m2: int, k2: int) -> float:
    return max(0.12, min(_ace_m_efficiency(m2), _ace_k_efficiency(k2)) * 0.95)


def _matmul_grid_candidates(m: int, n: int) -> list[dict[str, int]]:
    candidates = [
        {"rows": 4, "cols": 8},
        {"rows": 8, "cols": 4},
        {"rows": 2, "cols": 16},
        {"rows": 16, "cols": 2},
        {"rows": 1, "cols": 32},
        {"rows": 32, "cols": 1},
    ]
    scored: list[tuple[float, dict[str, int]]] = []
    aspect = m / max(1, n)
    for grid in candidates:
        tile_aspect = grid["rows"] / max(1, grid["cols"])
        score = abs(tile_aspect - aspect)
        if (grid["rows"], grid["cols"]) == (4, 8):
            score -= 0.05
        scored.append((score, grid))
    return [grid for _, grid in sorted(scored, key=lambda item: item[0])]


def _grid_virtualization_penalty(rows: int, cols: int) -> float:
    penalty = 1.0
    if cols > 8:
        penalty *= 1.0 + 0.12 * (cols - 8)
    if rows > 4:
        penalty *= 1.0 + 0.18 * (rows - 4)
    if (rows, cols) != (4, 8):
        penalty *= 1.05
    return penalty


def _matmul_ldm_bytes(m2: int, n2: int, k2: int, *, double_b: bool = True) -> int:
    a_bytes = m2 * k2 * 2
    b_bytes = k2 * n2 * 2
    c_bytes = m2 * n2 * 4
    return a_bytes + (2 * b_bytes if double_b else b_bytes) + c_bytes


def _matmul_k2_candidates() -> list[int]:
    return [32, 64, 96, 128, 160, 192, 256, 320, 384, 512, 640, 768, 1024]


def _matmul_m2_candidates() -> list[int]:
    return [16, 32, 64, 96, 128]


def _matmul_note_for_grid(rows: int, cols: int) -> tuple[list[str], str]:
    comm: list[str] = []
    if cols > 1:
        comm.append("A panel row-broadcast")
    else:
        comm.append("A local DMA only")
    if rows > 1:
        comm.append("B panel col-broadcast")
    else:
        comm.append("B local DMA only")
    if rows > 1 and cols > 1:
        reduction = "partial C stays local inside one cluster tile; no extra split-K reduction"
    else:
        reduction = "no inter-core reduction in the first-pass model"
    return comm, reduction


def _matmul_hardware_caps_for_schema(caps: dict[str, Any]) -> dict[str, float]:
    return {
        "ace_peak_tflops": float(caps["cluster_ace_peak_tflops"]),
        "dma_read_gbps": float(caps["cluster_dma_read_gbps"]),
        "dma_write_gbps": float(caps["cluster_dma_write_gbps"]),
        "rma_p2p_gbps": float(caps["cluster_rma_p2p_gbps"]),
        "rma_row_bcast_gbps": float(caps["cluster_rma_row_bcast_gbps"]),
        "rma_col_bcast_gbps": float(caps["cluster_rma_col_bcast_gbps"]),
        "ldm_bank_gbps": float(caps["ldm_bank_gbps"]),
        "ace_west_gbps": float(caps["ace_west_gbps"]),
        "ace_north_gbps": float(caps["ace_north_gbps"]),
        "ace_writeback_gbps": float(caps["ace_writeback_gbps"]),
    }


def _matmul_bottleneck_type(stage: str) -> str:
    if stage == "ace":
        return "compute"
    if stage in {"dma", "writeback"}:
        return "memory"
    return "mixed"


def _matmul_total_flops(m: int, n: int, k: int) -> int:
    return 2 * m * n * k


def _matmul_dma_segments(
    *,
    a_bytes: int,
    b_bytes: int,
    m2: int,
    n2: int,
    k2: int,
    rows: int,
    cols: int,
    extra_bytes: int = 0,
) -> list[dict[str, Any]]:
    segments: list[dict[str, Any]] = [
        {
            "name": "A",
            "bytes_total": a_bytes,
            "bytes_per_transfer": max(1, m2 * k2 * 2),
            "stride": rows > 1 or cols > 1,
            "hint": {"bm": max(1, m2), "bk": max(32, k2), "lda": max(32, k2) * max(1, cols)},
        },
        {
            "name": "B",
            "bytes_total": b_bytes,
            "bytes_per_transfer": max(1, n2 * k2 * 2),
            "stride": rows > 1,
            "hint": {"bm": max(1, n2), "bk": max(32, k2), "lda": max(32, k2) * max(1, rows)},
        },
    ]
    if extra_bytes > 0:
        segments.append(
            {
                "name": "workspace",
                "bytes_total": extra_bytes,
                "bytes_per_transfer": max(1, m2 * n2 * 4),
                "stride": False,
            }
        )
    return segments


def _matmul_dma_segments_pvc_original(
    *,
    a_bytes: int,
    b_bytes: int,
    m2: int,
    n2: int,
    k1: int,
    rows: int,
    cols: int,
) -> list[dict[str, Any]]:
    b_bk = max(32, k1 // max(1, rows))
    return [
        {
            "name": "A",
            "bytes_total": a_bytes,
            "bytes_per_transfer": max(1, m2 * k1 * 2),
            "stride": True,
            "hint": {"bm": max(1, m2), "bk": max(32, k1), "lda": max(32, k1) * max(1, cols)},
        },
        {
            "name": "B",
            "bytes_total": b_bytes,
            "bytes_per_transfer": max(1, n2 * b_bk * 2),
            "stride": True,
            "hint": {"bm": max(1, n2), "bk": b_bk, "lda": b_bk * max(1, rows)},
        },
    ]


def _matmul_feasible_region_summary(best: dict[str, Any], m: int, n: int, k: int) -> str:
    rows = int(best["grid"]["rows"])
    cols = int(best["grid"]["cols"])
    return (
        f"Selected from T1 GEMM candidates for locked shape M={m}, N={n}, K={k}. "
        f"Winning candidate family={best.get('family', 'traditional')}, grid={rows}x{cols}, "
        f"M1/N1/K1={best['m1']}/{best['n1']}/{best['k1']}, "
        f"M2/N2/K2={best['m2']}/{best['n2']}/{best['k2']}, "
        f"LDM={best['ldm_bytes']} bytes <= {DEFAULT_LDM_USABLE_BYTES} usable bytes, "
        f"estimated total={best['total_us']:.3f} us."
    )


def _matmul_candidate_traffic(best: dict[str, Any], m: int, n: int, k: int) -> dict[str, int]:
    compute_flops = _matmul_total_flops(m, n, k)
    if all(key in best for key in ("hbm_read_bytes", "hbm_write_bytes", "rma_bytes")):
        return {
            "hbm_read_bytes": int(best["hbm_read_bytes"]),
            "hbm_write_bytes": int(best["hbm_write_bytes"]),
            "rma_bytes": int(best["rma_bytes"]),
            "compute_flops": compute_flops,
        }

    m_tiles = _ceil_div(m, int(best["m1"]))
    n_tiles = _ceil_div(n, int(best["n1"]))
    k_tiles = _ceil_div(k, int(best["k1"]))
    read_per_k = int(best["m1"] * best["k1"] * 2 + best["k1"] * best["n1"] * 2)
    write_per_mn = int(best["m1"] * best["n1"] * 4)
    return {
        "hbm_read_bytes": m_tiles * n_tiles * k_tiles * read_per_k,
        "hbm_write_bytes": m_tiles * n_tiles * write_per_mn,
        "rma_bytes": int(best.get("rma_bytes_per_iter", 0)) * m_tiles * n_tiles * k_tiles,
        "compute_flops": compute_flops,
    }


def _fill_matmul_analysis_fields(
    out: dict[str, Any],
    *,
    best: dict[str, Any],
    m: int,
    n: int,
    k: int,
    candidate_notes: list[str],
) -> None:
    compute_flops = _matmul_total_flops(m, n, k)
    bottleneck_stage = str(best.get("bottleneck_stage", "dma"))
    bottleneck_type = _matmul_bottleneck_type(bottleneck_stage)

    op_def = out.setdefault("operator_definition", {})
    op_def["status"] = "defined"
    op_def["operator_family"] = "gemm"
    op_def.setdefault("summary", "")
    if not op_def.get("formula"):
        op_def["formula"] = "C[M,N] += A[M,K] * B[K,N]"
    if not op_def.get("io_tensors"):
        op_def["io_tensors"] = [
            {"name": "A", "dtype": "float16", "shape": ["M", "K"], "desc": "input matrix A"},
            {"name": "B", "dtype": "float16", "shape": ["K", "N"], "desc": "input matrix B"},
            {"name": "C", "dtype": "float32", "shape": ["M", "N"], "desc": "output matrix C"},
        ]
    if not op_def.get("kernel_apis"):
        op_def["kernel_apis"] = [
            "rt_dma_iget / rt_dma_iput for HBM-LDM movement",
            "rt_rma_iput / row/col communication for panel sharing or reduction",
            "rt_ace_load_west / rt_ace_writeback for fp16 GEMM on ACE",
        ]
    op_def["independent_axes"] = op_def.get("independent_axes") or ["M", "N"]
    op_def["reduction_axes"] = op_def.get("reduction_axes") or ["K"]
    op_def["reuse_axes"] = op_def.get("reuse_axes") or ["A reused across N tile", "B reused across M tile"]

    flops = out.setdefault("flops", {})
    flops.update(
        {
            "priority": "P0",
            "status": "defined",
            "formula_breakdown": [
                "Each C[m,n] accumulates K multiply-add operations.",
                f"flops_per_element = 2*K = {2 * k}",
                f"total_flops = 2*M*N*K = {compute_flops}",
            ],
            "flops_per_element": 2 * k,
            "total_flops": compute_flops,
            "notes": "Dense GEMM fp16 input with fp32 accumulation/output.",
        }
    )

    params = out.setdefault("parameters", {})
    params["status"] = "defined"
    params["shape_mode"] = "locked"
    params["shape_keys"] = ["M", "N", "K"]
    params["locked_shapes"] = {"M": m, "N": n, "K": k}
    params.setdefault("dtype_io", "float16")
    params.setdefault("dtype_acc", "float32")

    traffic = _matmul_candidate_traffic(best, m, n, k)
    visible_bytes = max(1, traffic["hbm_read_bytes"] + traffic["hbm_write_bytes"])
    stage = {
        "dma": round(float(best.get("t_dma", 0.0)), 6),
        "rma": round(float(best.get("t_rma", 0.0)), 6),
        "ace": round(float(best.get("t_ace", 0.0)), 6),
        "simd": round(float(best.get("t_simd", 0.0)), 6),
        "compute_group": round(max(float(best.get("t_ace", 0.0)), float(best.get("t_simd", 0.0))), 6),
        "writeback": round(float(best.get("t_wb", 0.0)), 6),
        "sync": round(float(best.get("t_sync", 0.0)), 6),
    }
    tp = out.setdefault("throughput_model", {})
    tp.update(
        {
            "priority": "P0",
            "status": "defined",
            "model_type": "layered_t1",
            "spe_clock_hz": int(T1_MATMUL_CAPS["spe_clock_hz"]),
            "peak_flops_per_spe": float(T1_MATMUL_CAPS["cluster_ace_peak_tflops"]) * 1e12 / 32,
            "hbm_bandwidth_gbps": float(T1_MATMUL_CAPS["cluster_dma_read_gbps"]),
            "arithmetic_intensity": round(compute_flops / visible_bytes, 6),
            "theoretical_peak_tflops": float(T1_MATMUL_CAPS["cluster_ace_peak_tflops"]),
            "upper_bound_tflops": round(float(best["tflops"]), 6),
            "hardware_caps": _matmul_hardware_caps_for_schema(T1_MATMUL_CAPS),
            "traffic_model": traffic,
            "stage_time_model_us": stage,
            "dma_calibration_segments": best.get("dma_calibration_segments", []),
            "constraint_discounts": [
                "ACE stage uses ACE.xlsx TFLOPS interpolation when available; otherwise falls back to T1 M2/K2 empirical envelope.",
                "DMA stage derives per-candidate A/B/workspace bm/bk/lda hints from the selected dataflow and uses DMA测试.xlsx exact/interpolated lookup when available; callers without hints fall back to nearest observed transfer shape, then aggregate T1 caps.",
                "SIMD stage uses SIMD使用手册-zh.md, simd-dual-issue-pattern.md, and 指令拍数和流水线派发.md: P0/P1 dual issue, vector FP add/mul/FMA around 5-cycle latency, FP16/FP32 convert around 2-cycle latency, local vector load/store around 5/4-cycle latency, with exposed cycles=max(P0 issue,P1 issue)+software-pipeline fill/drain.",
                "ACE and SIMD are modeled as an overlapped compute group: compute_group=max(T_ace,T_simd), so SIMD epilogue/reduce work can be hidden by ACE when pipeline dependencies allow.",
                "RMA bandwidth uses row/col/p2p caps according to selected candidate communication pattern.",
                "LDM feasibility uses usable 235KB budget and first-pass bank split.",
                *_load_calibration_tables().get("notes", []),
            ],
            "achievable_throughput_note": (
                f"Selected candidate estimates {best['total_us']:.3f} us total and "
                f"{best['tflops']:.3f} TFLOPS for the locked shape. This is a design-time "
                "upper bound before profiler calibration."
            ),
            "bottleneck_type": bottleneck_type,
            "bottleneck_analysis": (
                f"Per steady-state candidate stage times are dma={stage['dma']} us, "
                f"rma={stage['rma']} us, ace={stage['ace']} us, simd={stage['simd']} us, "
                f"compute_group={stage['compute_group']} us, "
                f"writeback={stage['writeback']} us, sync={stage['sync']} us. "
                f"The max exposed stage is {bottleneck_stage}, so the first-pass "
                f"bottleneck classification is {bottleneck_type}."
            ),
            "target_efficiency": round(
                min(1.0, float(best["tflops"]) / float(T1_MATMUL_CAPS["cluster_ace_peak_tflops"])),
                6,
            ),
        }
    )

    tile = out.setdefault("tiling_strategy", {})
    search_space = tile.setdefault("search_space", {})
    search_space.setdefault("notes", candidate_notes)
    tile["feasible_region_summary"] = _matmul_feasible_region_summary(best, m, n, k)
    out["status"] = "review"


def deep_merge(base: dict[str, Any], patch: dict[str, Any]) -> dict[str, Any]:
    out = deepcopy(base)
    for key, value in patch.items():
        if isinstance(value, dict) and isinstance(out.get(key), dict):
            out[key] = deep_merge(out[key], value)
        else:
            out[key] = value
    return out


def _section_has_content(section: dict[str, Any]) -> bool:
    if not section:
        return False
    for key, value in section.items():
        if key in ("status", "priority"):
            continue
        if value in (None, "", [], {}):
            continue
        return True
    return False


def normalize_design(doc: dict[str, Any]) -> dict[str, Any]:
    """Fill missing optional fields so validate/render/scaffold work on partial JSON."""
    out = deepcopy(doc)
    ctype = out.get("operator_definition", {}).get("computation_type", "elementwise")

    if out.get("status") not in VALID_DOC_STATUS:
        out["status"] = "draft"

    wf = out.setdefault("workflow", {})
    if wf.get("current_focus") not in VALID_WORKFLOW_FOCUS:
        wf["current_focus"] = "operator_definition"
    wf.setdefault("mode", "iterative")
    wf.setdefault("description", "")
    wf.setdefault("iteration_round", 1)
    wf.setdefault("history", [])

    for section_key, defaults in (
        ("operator_definition", {"priority": "P0", "status": "draft"}),
        ("flops", {"priority": "P0", "status": "draft"}),
        ("parameters", {"priority": "P0", "status": "draft"}),
        ("throughput_model", {"priority": "P0", "status": "draft"}),
        ("ldm_planning", {"priority": "P0", "status": "draft"}),
    ):
        section = out.setdefault(section_key, {})
        for k, v in defaults.items():
            section.setdefault(k, v)
        if section.get("status") not in VALID_SECTION_STATUS:
            section["status"] = "defined" if _section_has_content(section) else "draft"

    for section_key in SECTIONS_REQUIRING_STATUS:
        section = out.setdefault(section_key, {})
        if section.get("status") not in VALID_SECTION_STATUS:
            section["status"] = "defined" if _section_has_content(section) else "draft"

    scaffold = out.setdefault("scaffold", {})
    if scaffold.get("family") not in VALID_SCAFFOLD_FAMILIES:
        family = COMPUTATION_TYPE_TO_SCAFFOLD.get(str(ctype), "elewise")
        scaffold["family"] = family
    scaffold.setdefault("smoke", True)
    scaffold.setdefault("patch_runsh", True)
    scaffold.setdefault("upgrade_common", True)

    out.setdefault("golden", {"framework": "pytorch", "atol": 1e-5, "rtol": 3e-2})

    op_def = out.setdefault("operator_definition", {})
    op_def.setdefault("operator_family", _default_operator_family(str(ctype)))
    op_def.setdefault("independent_axes", [])
    op_def.setdefault("reduction_axes", [])
    op_def.setdefault("reuse_axes", [])

    params = out.setdefault("parameters", {})
    params.setdefault("shape_mode", "locked")
    shapes = params.setdefault("shapes", {})
    if "smoke" not in shapes:
        shapes["smoke"] = _default_shapes(str(ctype), smoke=True)
    if "default" not in shapes:
        shapes["default"] = _default_shapes(str(ctype), smoke=False)
    params.setdefault("shape_keys", list(shapes.get("default", {}).keys()))
    params.setdefault("locked_shapes", deepcopy(shapes.get("default", {})))
    params.setdefault("shape_ranges", {})
    params.setdefault(
        "ideal_shape_goal",
        {
            "enabled": False,
            "optimize_for": "throughput",
            "fixed_axes": {},
            "search_axes": [],
            "notes": "",
        },
    )

    ldm = out.setdefault("ldm_planning", {})
    ldm.setdefault("budget_bytes_per_spe", DEFAULT_LDM_BUDGET_BYTES)
    ldm.setdefault("usable_budget_bytes", DEFAULT_LDM_USABLE_BYTES)
    ldm.setdefault("buffers", [])
    if ldm.get("total_bytes") in (None, 0) and ldm.get("buffers"):
        ldm["total_bytes"] = sum(int(b.get("bytes", 0)) for b in ldm["buffers"])
    ldm.setdefault("fits_budget", True)
    ldm.setdefault("bank_allocation", [])
    ldm.setdefault("pressure_summary", "")
    ldm.setdefault("reuse_notes", "")

    tp = out.setdefault("throughput_model", {})
    tp.setdefault("model_type", "layered_t1")
    tp.setdefault(
        "hardware_caps",
        {
            "ace_peak_tflops": None,
            "dma_read_gbps": None,
            "dma_write_gbps": None,
            "rma_p2p_gbps": None,
            "rma_row_bcast_gbps": None,
            "rma_col_bcast_gbps": None,
            "ldm_bank_gbps": None,
            "ace_west_gbps": None,
            "ace_north_gbps": None,
            "ace_writeback_gbps": None,
        },
    )
    tp.setdefault(
        "traffic_model",
        {
            "hbm_read_bytes": None,
            "hbm_write_bytes": None,
            "rma_bytes": None,
            "compute_flops": None,
        },
    )
    tp.setdefault(
        "stage_time_model_us",
        {"dma": None, "rma": None, "ace": None, "writeback": None, "sync": None},
    )
    tp.setdefault("constraint_discounts", [])
    tp.setdefault("upper_bound_tflops", None)

    tile = out.setdefault("tiling_strategy", {})
    tile.setdefault("hierarchy", [])
    tile.setdefault("selected_tile", {})
    tile.setdefault("search_space", {})
    tile.setdefault("constraints", [])
    tile.setdefault("feasible_region_summary", "")

    core = out.setdefault("core_partition_strategy", {})
    core.setdefault("grid_shape", {})
    core.setdefault("split_dimensions", [])
    core.setdefault("communication_pattern", [])
    core.setdefault("reduction_pattern", "")
    core.setdefault("loader_assignment", [])

    dbuf = out.setdefault("double_buffer_strategy", {})
    dbuf.setdefault("stage_buffers", [])
    dbuf.setdefault("enable_condition", "")
    dbuf.setdefault("blocked_by", [])

    pipe = out.setdefault("pipeline_layout", {})
    pipe.setdefault("steady_state_stages", [])
    pipe.setdefault("overlap_pairs", [])
    pipe.setdefault("wait_points", [])
    pipe.setdefault("fill_drain_notes", "")

    return out


def is_empty_value(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and not value.strip():
        return True
    if isinstance(value, (list, dict)) and len(value) == 0:
        return True
    return False


def _require_non_empty(
    errors: list[str],
    path: str,
    value: Any,
    *,
    allow_false: bool = False,
) -> None:
    if allow_false and value is False:
        return
    if is_empty_value(value):
        errors.append(f"[必填] {path} 不得为空（null / 空字符串 / 空数组 / 空对象）")


def _require_section_defined(errors: list[str], section_path: str, section: dict[str, Any]) -> None:
    status = section.get("status")
    if status != "defined":
        errors.append(f"[必填] {section_path}.status 必须为 defined（当前: {status!r}）")


def _require_shape_block(
    errors: list[str],
    path: str,
    shapes: dict[str, Any],
    *,
    shape_keys: list[str] | None = None,
) -> None:
    keys = shape_keys if shape_keys is not None else list(shapes.keys())
    if not keys:
        errors.append(f"[必填] {path} 不得为空对象")
        return
    for key in keys:
        _require_non_empty(errors, f"{path}.{key}", shapes.get(key))


def _require_dict_keys(
    errors: list[str],
    path: str,
    data: dict[str, Any],
    keys: list[str],
    *,
    allow_false_keys: frozenset[str] = frozenset(),
) -> None:
    _require_non_empty(errors, path, data)
    for key in keys:
        _require_non_empty(
            errors,
            f"{path}.{key}",
            data.get(key),
            allow_false=key in allow_false_keys,
        )


def _resolve_shape_keys(doc: dict[str, Any]) -> list[str]:
    """Determine expected shape keys from requirement.json or computation_type."""
    shape_keys = doc.get("parameters", {}).get("shape_keys", [])
    if shape_keys:
        return list(shape_keys)
    req_path = doc.get("requirement_path")
    if req_path:
        try:
            req = json.loads(Path(req_path).read_text(encoding="utf-8"))
            keys = req.get("shape_keys", [])
            if keys:
                return keys
        except (OSError, json.JSONDecodeError):
            pass
    ctype = doc.get("operator_definition", {}).get("computation_type", "")
    if ctype == "matmul":
        return ["M", "N", "K"]
    params = doc.get("parameters", {})
    locked = params.get("locked_shapes", {})
    if locked:
        return list(locked.keys())
    smoke = params.get("shapes", {}).get("smoke", {})
    if smoke:
        return list(smoke.keys())
    return ["N", "H", "W", "C"]


def _render_shapes(shapes: dict[str, Any]) -> str:
    """Render shape dicts generically, e.g. 'M=512 N=256 K=8192'."""
    lines = []
    for mode in ("smoke", "default", "locked_shapes", "shape_ranges"):
        mode_shapes = shapes.get(mode, {})
        if mode_shapes:
            parts = " ".join(f"{k}={v}" for k, v in mode_shapes.items())
            lines.append(f"**{mode}**：{parts}")
    return "\n\n".join(lines)


def validate_design_completeness(doc: dict[str, Any]) -> list[str]:
    """Strict gate: every schema field must be filled before render-md / scaffold."""
    errors: list[str] = []

    _require_non_empty(errors, "title", doc.get("title"))
    status = doc.get("status")
    if status in (None, "draft"):
        errors.append("[必填] status 不得为 draft；完成设计后设为 review 或 approved")
    if doc.get("workflow", {}).get("current_focus") != "complete":
        errors.append("[必填] workflow.current_focus 必须为 complete")

    wf = doc.get("workflow", {})
    _require_non_empty(errors, "workflow.mode", wf.get("mode"))
    _require_non_empty(errors, "workflow.description", wf.get("description"))
    _require_non_empty(errors, "workflow.iteration_round", wf.get("iteration_round"))
    _require_non_empty(errors, "workflow.history", wf.get("history"))
    for idx, item in enumerate(wf.get("history", [])):
        _require_non_empty(errors, f"workflow.history[{idx}]", item)

    op_def = doc.get("operator_definition", {})
    _require_section_defined(errors, "operator_definition", op_def)
    _require_non_empty(errors, "operator_definition.summary", op_def.get("summary"))
    _require_non_empty(errors, "operator_definition.operator_family", op_def.get("operator_family"))
    _require_non_empty(errors, "operator_definition.formula", op_def.get("formula"))
    _require_non_empty(errors, "operator_definition.logic_steps", op_def.get("logic_steps"))
    for idx, step in enumerate(op_def.get("logic_steps", [])):
        _require_non_empty(errors, f"operator_definition.logic_steps[{idx}]", step)
    io_tensors = op_def.get("io_tensors", [])
    _require_non_empty(errors, "operator_definition.io_tensors", io_tensors)
    for idx, tensor in enumerate(io_tensors):
        for key in ("name", "dtype", "shape", "desc"):
            _require_non_empty(errors, f"operator_definition.io_tensors[{idx}].{key}", tensor.get(key))
    _require_non_empty(errors, "operator_definition.kernel_apis", op_def.get("kernel_apis"))
    for idx, api in enumerate(op_def.get("kernel_apis", [])):
        _require_non_empty(errors, f"operator_definition.kernel_apis[{idx}]", api)
    _require_non_empty(errors, "operator_definition.independent_axes", op_def.get("independent_axes"))
    _require_non_empty(errors, "operator_definition.reduction_axes", op_def.get("reduction_axes"))
    _require_non_empty(errors, "operator_definition.reuse_axes", op_def.get("reuse_axes"))

    flops = doc.get("flops", {})
    _require_section_defined(errors, "flops", flops)
    _require_non_empty(errors, "flops.formula_breakdown", flops.get("formula_breakdown"))
    for idx, item in enumerate(flops.get("formula_breakdown", [])):
        _require_non_empty(errors, f"flops.formula_breakdown[{idx}]", item)
    _require_non_empty(errors, "flops.flops_per_element", flops.get("flops_per_element"))
    if is_empty_value(flops.get("total_flops")) and is_empty_value(flops.get("notes")):
        errors.append("[必填] flops.total_flops 与 flops.notes 至少填写一项")
    _require_non_empty(errors, "flops.notes", flops.get("notes"))

    params = doc.get("parameters", {})
    _require_section_defined(errors, "parameters", params)
    _require_non_empty(errors, "parameters.shape_mode", params.get("shape_mode"))
    _require_non_empty(errors, "parameters.shape_keys", params.get("shape_keys"))
    _require_non_empty(errors, "parameters.dtype_io", params.get("dtype_io"))
    _require_non_empty(errors, "parameters.dtype_acc", params.get("dtype_acc"))
    shapes = params.get("shapes", {})
    _require_non_empty(errors, "parameters.shapes.smoke", shapes.get("smoke"))
    _require_non_empty(errors, "parameters.shapes.default", shapes.get("default"))
    design_shape_keys = _resolve_shape_keys(doc)
    _require_shape_block(errors, "parameters.shapes.smoke", shapes.get("smoke", {}), shape_keys=design_shape_keys)
    _require_shape_block(errors, "parameters.shapes.default", shapes.get("default", {}), shape_keys=design_shape_keys)
    shape_mode = params.get("shape_mode")
    if shape_mode == "locked":
        _require_shape_block(errors, "parameters.locked_shapes", params.get("locked_shapes", {}), shape_keys=design_shape_keys)
    elif shape_mode == "ranged":
        _require_non_empty(errors, "parameters.shape_ranges", params.get("shape_ranges"))
    elif shape_mode == "ideal_search":
        goal = params.get("ideal_shape_goal", {})
        _require_dict_keys(
            errors,
            "parameters.ideal_shape_goal",
            goal,
            ["enabled", "optimize_for", "fixed_axes", "search_axes", "notes"],
            allow_false_keys=frozenset({"enabled"}),
        )
    hyper = params.get("hyperparams", {})
    if op_def.get("computation_type") == "normalization":
        _require_non_empty(errors, "parameters.hyperparams.eps", hyper.get("eps"))

    tp = doc.get("throughput_model", {})
    _require_section_defined(errors, "throughput_model", tp)
    _require_non_empty(errors, "throughput_model.model_type", tp.get("model_type"))
    for key in (
        "spe_clock_hz",
        "peak_flops_per_spe",
        "hbm_bandwidth_gbps",
        "arithmetic_intensity",
        "theoretical_peak_tflops",
        "upper_bound_tflops",
        "achievable_throughput_note",
        "bottleneck_analysis",
        "target_efficiency",
    ):
        _require_non_empty(errors, f"throughput_model.{key}", tp.get(key))
    _require_dict_keys(
        errors,
        "throughput_model.hardware_caps",
        tp.get("hardware_caps", {}),
        [
            "ace_peak_tflops",
            "dma_read_gbps",
            "dma_write_gbps",
            "rma_p2p_gbps",
            "rma_row_bcast_gbps",
            "rma_col_bcast_gbps",
            "ldm_bank_gbps",
            "ace_west_gbps",
            "ace_north_gbps",
            "ace_writeback_gbps",
        ],
    )
    _require_dict_keys(
        errors,
        "throughput_model.traffic_model",
        tp.get("traffic_model", {}),
        ["hbm_read_bytes", "hbm_write_bytes", "rma_bytes", "compute_flops"],
    )
    _require_dict_keys(
        errors,
        "throughput_model.stage_time_model_us",
        tp.get("stage_time_model_us", {}),
        ["dma", "rma", "ace", "writeback", "sync"],
    )
    _require_non_empty(errors, "throughput_model.constraint_discounts", tp.get("constraint_discounts"))
    bottleneck = tp.get("bottleneck_type")
    if bottleneck not in STRICT_BOTTLENECK_TYPES:
        errors.append(
            f"[必填] throughput_model.bottleneck_type 不得为 unknown（当前: {bottleneck!r}）"
        )

    tile = doc.get("tiling_strategy", {})
    _require_section_defined(errors, "tiling_strategy", tile)
    for key in (
        "tile_axis",
        "tile_elems",
        "dma_bytes_per_tile",
        "alignment_bytes",
        "hierarchy",
        "selected_tile",
        "search_space",
        "constraints",
        "feasible_region_summary",
        "rationale",
    ):
        _require_non_empty(errors, f"tiling_strategy.{key}", tile.get(key))

    core = doc.get("core_partition_strategy", {})
    _require_section_defined(errors, "core_partition_strategy", core)
    for key in (
        "slave_core_num",
        "partition_axis",
        "strategy",
        "pattern",
        "grid_shape",
        "split_dimensions",
        "communication_pattern",
        "reduction_pattern",
        "loader_assignment",
        "rationale",
    ):
        _require_non_empty(errors, f"core_partition_strategy.{key}", core.get(key))

    dbuf = doc.get("double_buffer_strategy", {})
    _require_section_defined(errors, "double_buffer_strategy", dbuf)
    _require_non_empty(errors, "double_buffer_strategy.enabled", dbuf.get("enabled"), allow_false=True)
    _require_non_empty(errors, "double_buffer_strategy.buffered_tensors", dbuf.get("buffered_tensors"))
    _require_non_empty(errors, "double_buffer_strategy.buffer_depth", dbuf.get("buffer_depth"))
    _require_non_empty(
        errors,
        "double_buffer_strategy.overlap_dma_compute",
        dbuf.get("overlap_dma_compute"),
        allow_false=True,
    )
    _require_non_empty(errors, "double_buffer_strategy.stage_buffers", dbuf.get("stage_buffers"))
    _require_non_empty(errors, "double_buffer_strategy.enable_condition", dbuf.get("enable_condition"))
    _require_non_empty(errors, "double_buffer_strategy.blocked_by", dbuf.get("blocked_by"))
    _require_non_empty(errors, "double_buffer_strategy.rationale", dbuf.get("rationale"))

    ldm = doc.get("ldm_planning", {})
    _require_section_defined(errors, "ldm_planning", ldm)
    _require_non_empty(errors, "ldm_planning.budget_bytes_per_spe", ldm.get("budget_bytes_per_spe"))
    _require_non_empty(errors, "ldm_planning.usable_budget_bytes", ldm.get("usable_budget_bytes"))
    buffers = ldm.get("buffers", [])
    _require_non_empty(errors, "ldm_planning.buffers", buffers)
    for idx, buf in enumerate(buffers):
        for key in ("name", "dtype", "elems", "bytes", "purpose"):
            _require_non_empty(errors, f"ldm_planning.buffers[{idx}].{key}", buf.get(key))
    _require_non_empty(errors, "ldm_planning.total_bytes", ldm.get("total_bytes"))
    _require_non_empty(errors, "ldm_planning.fits_budget", ldm.get("fits_budget"), allow_false=True)
    _require_non_empty(errors, "ldm_planning.bank_allocation", ldm.get("bank_allocation"))
    _require_non_empty(errors, "ldm_planning.pressure_summary", ldm.get("pressure_summary"))
    _require_non_empty(errors, "ldm_planning.reuse_notes", ldm.get("reuse_notes"))
    _require_non_empty(errors, "ldm_planning.allocation_notes", ldm.get("allocation_notes"))

    pipe = doc.get("pipeline_layout", {})
    _require_section_defined(errors, "pipeline_layout", pipe)
    _require_non_empty(errors, "pipeline_layout.dataflow_description", pipe.get("dataflow_description"))
    _require_non_empty(errors, "pipeline_layout.passes", pipe.get("passes"))
    for idx, item in enumerate(pipe.get("passes", [])):
        _require_non_empty(errors, f"pipeline_layout.passes[{idx}]", item)
    _require_non_empty(errors, "pipeline_layout.sync_points", pipe.get("sync_points"))
    for idx, item in enumerate(pipe.get("sync_points", [])):
        _require_non_empty(errors, f"pipeline_layout.sync_points[{idx}]", item)
    _require_non_empty(errors, "pipeline_layout.steady_state_stages", pipe.get("steady_state_stages"))
    _require_non_empty(errors, "pipeline_layout.overlap_pairs", pipe.get("overlap_pairs"))
    _require_non_empty(errors, "pipeline_layout.wait_points", pipe.get("wait_points"))
    _require_non_empty(errors, "pipeline_layout.fill_drain_notes", pipe.get("fill_drain_notes"))

    scaffold = doc.get("scaffold", {})
    for key in ("family", "smoke", "patch_runsh", "upgrade_common"):
        _require_non_empty(errors, f"scaffold.{key}", scaffold.get(key), allow_false=True)

    golden = doc.get("golden", {})
    for key in ("framework", "atol", "rtol"):
        _require_non_empty(errors, f"golden.{key}", golden.get(key))

    for ext_key, ext_val in doc.items():
        if ext_key in STANDARD_TOP_LEVEL_KEYS:
            continue
        if isinstance(ext_val, dict):
            for sub_key, sub_val in ext_val.items():
                _require_non_empty(errors, f"{ext_key}.{sub_key}", sub_val)
        else:
            _require_non_empty(errors, ext_key, ext_val)

    return errors


def validate_design(doc: dict[str, Any], *, strict: bool = False) -> list[str]:
    errors: list[str] = []
    if doc.get("schema_version") != SCHEMA_VERSION:
        errors.append(f"unsupported schema_version (expected {SCHEMA_VERSION})")

    try:
        validate_op_name(str(doc.get("op", "")))
    except ValueError as exc:
        errors.append(str(exc))

    op_def = doc.get("operator_definition", {})
    ctype = op_def.get("computation_type", "")
    if ctype not in VALID_COMPUTATION_TYPES:
        errors.append(f"operator_definition.computation_type invalid: {ctype}")
    if is_empty_value(op_def.get("operator_family")):
        errors.append("operator_definition.operator_family missing")

    params = doc.get("parameters", {})
    shape_mode = params.get("shape_mode", "")
    if shape_mode and shape_mode not in {"locked", "ranged", "ideal_search"}:
        errors.append(f"parameters.shape_mode invalid: {shape_mode}")
    smoke = params.get("shapes", {}).get("smoke", {})
    shape_keys = _resolve_shape_keys(doc)
    for key in shape_keys:
        if key not in smoke:
            errors.append(f"parameters.shapes.smoke missing {key}")
    if ctype == "matmul" and shape_keys != ["M", "N", "K"]:
        errors.append("matmul design must use shape_keys ['M', 'N', 'K']")

    tp = doc.get("throughput_model", {})
    if tp.get("bottleneck_type") not in VALID_BOTTLENECK_TYPES:
        errors.append("throughput_model.bottleneck_type invalid")
    stage_model = tp.get("stage_time_model_us", {})
    if stage_model:
        filled = [k for k, v in stage_model.items() if not is_empty_value(v)]
        if len(filled) and len(filled) != len(stage_model):
            errors.append("throughput_model.stage_time_model_us must be fully specified")

    ldm = doc.get("ldm_planning", {})
    budget = int(ldm.get("budget_bytes_per_spe", 0))
    usable_budget = int(ldm.get("usable_budget_bytes", 0))
    total = int(ldm.get("total_bytes", 0))
    if budget > 0 and total > budget:
        errors.append(f"ldm_planning.total_bytes ({total}) exceeds budget ({budget})")
    if usable_budget > 0 and total > usable_budget:
        errors.append(
            f"ldm_planning.total_bytes ({total}) exceeds usable_budget_bytes ({usable_budget})"
        )
    if usable_budget and budget and usable_budget > budget:
        errors.append("ldm_planning.usable_budget_bytes cannot exceed budget_bytes_per_spe")

    core = doc.get("core_partition_strategy", {})
    grid = core.get("grid_shape", {})
    rows = grid.get("rows")
    cols = grid.get("cols")
    cores = core.get("slave_core_num")
    if rows and cols and cores and int(rows) * int(cols) != int(cores):
        errors.append("core_partition_strategy.grid_shape rows*cols must equal slave_core_num")

    tile = doc.get("tiling_strategy", {})
    selected_tile = tile.get("selected_tile", {})
    if ctype == "matmul" and selected_tile:
        for key in ("M1", "N1", "K1", "M2", "N2", "K2"):
            if key not in selected_tile:
                errors.append(f"tiling_strategy.selected_tile missing {key} for matmul")
        if "K2" in selected_tile and int(selected_tile["K2"]) % 32 != 0:
            errors.append("tiling_strategy.selected_tile.K2 must be 32-aligned on T1")
        if "N2" in selected_tile and int(selected_tile["N2"]) % 32 != 0:
            errors.append("tiling_strategy.selected_tile.N2 must be 32-aligned on T1")

    dbuf = doc.get("double_buffer_strategy", {})
    if dbuf.get("enabled") and not dbuf.get("stage_buffers"):
        errors.append("double_buffer_strategy.enabled=true requires stage_buffers")

    pipe = doc.get("pipeline_layout", {})
    if dbuf.get("enabled") and not pipe.get("overlap_pairs"):
        errors.append("pipeline_layout.overlap_pairs required when double buffering is enabled")

    for section in SECTIONS_REQUIRING_STATUS:
        st = doc.get(section, {}).get("status")
        if st is not None and st not in VALID_SECTION_STATUS:
            errors.append(f"{section}.status must be draft or defined")

    scaffold = doc.get("scaffold", {})
    sf = scaffold.get("family", "")
    if sf and sf not in VALID_SCAFFOLD_FAMILIES:
        errors.append(f"scaffold.family must be one of {sorted(VALID_SCAFFOLD_FAMILIES)}")

    focus = doc.get("workflow", {}).get("current_focus", "")
    if focus and focus not in VALID_WORKFLOW_FOCUS:
        errors.append(f"workflow.current_focus invalid: {focus}")

    if doc.get("status") and doc.get("status") not in VALID_DOC_STATUS:
        errors.append(f"status must be one of {sorted(VALID_DOC_STATUS)}")

    if strict:
        errors.extend(validate_design_completeness(doc))

    return errors


def list_incomplete_fields(doc: dict[str, Any]) -> list[str]:
    """Return field paths that fail completeness (for Agent iteration)."""
    return validate_design_completeness(doc)


def load_design(path: Path, *, normalize: bool = True, strict: bool = True) -> dict[str, Any]:
    if not path.is_file():
        raise FileNotFoundError(f"design not found: {path}")
    doc = json.loads(path.read_text(encoding="utf-8"))
    if normalize:
        doc = normalize_design(doc)
    errors = validate_design(doc, strict=strict)
    if errors:
        raise ValueError("design validation failed:\n  - " + "\n  - ".join(errors))
    return doc


def resolve_design_path(demo_root: Path, op: str, explicit: str | None = None) -> Path:
    if explicit:
        path = Path(explicit).expanduser().resolve()
        if not path.is_file():
            raise FileNotFoundError(f"design file not found: {path}")
        return path
    default = design_paths(demo_root, op).json_path
    if not default.is_file():
        raise FileNotFoundError(
            f"no design at {default}; run init_design.py --op {op} first"
        )
    return default


def shapes_to_context(shapes: dict[str, int]) -> dict[str, str]:
    return {f"k{k}": str(v) for k, v in shapes.items()}


def infer_scaffold_family(doc: dict[str, Any]) -> str:
    scaffold = doc.get("scaffold", {})
    if scaffold.get("family"):
        return str(scaffold["family"])
    ctype = doc.get("operator_definition", {}).get("computation_type", "elementwise")
    family = COMPUTATION_TYPE_TO_SCAFFOLD.get(str(ctype))
    if not family:
        raise ValueError(
            f"computation_type '{ctype}' has no scaffold mapping; set scaffold.family explicitly"
        )
    return family


def apply_design_to_scaffold(doc: dict[str, Any], *, smoke: bool | None = None) -> dict[str, Any]:
    family = infer_scaffold_family(doc)
    scaffold = doc.get("scaffold", {})
    use_smoke = bool(scaffold.get("smoke", True)) if smoke is None else smoke
    shape_key = "smoke" if use_smoke else "default"
    shapes = doc["parameters"]["shapes"][shape_key]
    hyper = doc.get("parameters", {}).get("hyperparams", {})
    golden = doc.get("golden", {})
    return {
        "family": family,
        "smoke": use_smoke,
        "shapes": shapes,
        "patch_runsh": bool(scaffold.get("patch_runsh", True)),
        "upgrade_common": bool(scaffold.get("upgrade_common", True)),
        "eps": str(hyper.get("eps", "1e-5")),
        "slave_core_num": int(
            doc.get("core_partition_strategy", {}).get("slave_core_num", 32)
        ),
        "golden_atol": golden.get("atol"),
        "golden_rtol": golden.get("rtol"),
        "design_path": None,
    }


def _bullets(items: list[str]) -> str:
    if not items:
        return "- （待填写）\n"
    return "".join(f"- {item}\n" for item in items)


def _section_status_badge(section: dict[str, Any]) -> str:
    priority = section.get("priority", "")
    status = section.get("status", "draft")
    p = f" **{priority}**" if priority else ""
    return f"`{status}`{p}"


def _kv_table(data: dict[str, Any], *, skip: frozenset[str] = frozenset()) -> str:
    rows = ""
    for key, value in data.items():
        if key in skip:
            continue
        if isinstance(value, (list, dict)):
            rendered = json.dumps(value, ensure_ascii=False)
        else:
            rendered = value
        rows += f"| {key} | {rendered} |\n"
    return rows or "| （无） | |\n"


def _render_buffer_table(buffers: list[dict[str, Any]]) -> str:
    if not buffers:
        return "| （待填写） |\n"
    col_order = ["name", "dtype", "elems", "bytes", "purpose"]
    cols: list[str] = []
    for name in col_order:
        if any(name in b for b in buffers):
            cols.append(name)
    for b in buffers:
        for key in b:
            if key not in cols:
                cols.append(key)
    header = "| " + " | ".join(cols) + " |"
    sep = "| " + " | ".join("---" for _ in cols) + " |"
    body = "\n".join(
        "| " + " | ".join(str(b.get(c, "")) for c in cols) + " |" for b in buffers
    )
    return f"{header}\n{sep}\n{body}\n"


def _render_extensions(doc: dict[str, Any]) -> str:
    extras = {k: v for k, v in doc.items() if k not in STANDARD_TOP_LEVEL_KEYS}
    if not extras:
        return ""
    blocks = ""
    for key, value in extras.items():
        title = key.replace("_", " ").title()
        if isinstance(value, dict):
            blocks += f"\n### {title}\n\n| 字段 | 值 |\n|------|-----|\n{_kv_table(value)}"
        else:
            blocks += f"\n### {title}\n\n{json.dumps(value, ensure_ascii=False, indent=2)}\n"
    return f"""
---

## 9. 扩展字段（design.json 额外键）

{blocks}
"""


def render_design_md(doc: dict[str, Any]) -> str:
    op = doc["op"]
    wf = doc.get("workflow", {})
    op_def = doc.get("operator_definition", {})
    flops = doc.get("flops", {})
    params = doc.get("parameters", {})
    tp = doc.get("throughput_model", {})
    tile = doc.get("tiling_strategy", {})
    core = doc.get("core_partition_strategy", {})
    dbuf = doc.get("double_buffer_strategy", {})
    ldm = doc.get("ldm_planning", {})
    pipe = doc.get("pipeline_layout", {})
    scaffold = doc.get("scaffold", {})
    golden = doc.get("golden", {})
    shapes = params.get("shapes", {})
    shape_overview = {
        "smoke": shapes.get("smoke", {}),
        "default": shapes.get("default", {}),
        "locked_shapes": params.get("locked_shapes", {}),
        "shape_ranges": params.get("shape_ranges", {}),
    }

    workflow_steps = "\n".join(
        f"{label}{' **[P0]**' if p0 else ''}" for _, label, p0 in WORKFLOW_ORDER
    )

    history_block = ""
    history = wf.get("history", [])
    if history:
        history_block = "\n### 迭代历史\n\n" + _bullets(history)

    return f"""# 算子设计：{doc.get('title', op)}

> **算子**：`{op}` · **状态**：{doc.get('status', 'draft')} · **schema**：{doc.get('schema_version')}
> **创建**：{doc.get('created_at', '')} · **更新**：{doc.get('updated_at', '')}
> **迭代轮次**：{wf.get('iteration_round', 1)} · **当前焦点**：{wf.get('current_focus', '')}

机器可读：`design.json` -> teco-ops interface + UAL 落地

---

## 设计流程（循环迭代，非单链路）

{wf.get('description', '')}
{history_block}

```
{workflow_steps}
    ↓
输出：算子设计文档（本文件 + design.json）
```

> **P0（优先做）**：算子定义 + FLOPS、参数 + 吞吐建模、LDM 规划
> 分块 / 分核 / 双缓冲 / 流水可迭代细化，逼近理论峰值。

---

## 1. 算子定义 {_section_status_badge(op_def)}

**摘要**：{op_def.get('summary') or '（待填写）'}

| 字段 | 值 |
|------|-----|
| computation_type | `{op_def.get('computation_type')}` |
| operator_family | `{op_def.get('operator_family')}` |
| formula | 见下 |

```
{op_def.get('formula') or '（待填写）'}
```

### 计算步骤

{_bullets(op_def.get('logic_steps', []))}

### IO

{_bullets([f"{t.get('name')}: {t.get('dtype')} {t.get('shape')} — {t.get('desc', '')}" for t in op_def.get('io_tensors', [])])}

### KERNEL API

{_bullets(op_def.get('kernel_apis', []))}

---

## 1b. 计算 FLOPS {_section_status_badge(flops)}

| 项 | 值 |
|----|-----|
| flops_per_element | {flops.get('flops_per_element')} |
| total_flops | {flops.get('total_flops')} |

### 公式拆解

{_bullets(flops.get('formula_breakdown', []))}

**备注**：{flops.get('notes', '') or '（无）'}

---

## 2. 参数 {_section_status_badge(params)}

| 字段 | 值 |
|------|-----|
| shape_mode | `{params.get('shape_mode')}` |
| dtype_io | `{params.get('dtype_io')}` |
| dtype_acc | `{params.get('dtype_acc')}` |

{_render_shapes(shape_overview)}

**超参**：`{json.dumps(params.get('hyperparams', {}), ensure_ascii=False)}`

**ideal_shape_goal**：`{json.dumps(params.get('ideal_shape_goal', {}), ensure_ascii=False)}`

---

## 2b. 吞吐建模与瓶颈 {_section_status_badge(tp)}

| 项 | 值 |
|----|-----|
{_kv_table(tp, skip=frozenset({"priority", "status", "bottleneck_analysis", "achievable_throughput_note"}))}
**瓶颈分析**：{tp.get('bottleneck_analysis') or '（待填写）'}

**可达吞吐说明**：{tp.get('achievable_throughput_note') or '（待填写）'}

---

## 3. 分块策略 {_section_status_badge(tile)}

| 项 | 值 |
|----|-----|
{_kv_table(tile, skip=frozenset({"status", "priority", "rationale"}))}
**依据**：{tile.get('rationale') or '（待填写）'}

---

## 4. 分核策略 {_section_status_badge(core)}

| 项 | 值 |
|----|-----|
{_kv_table(core, skip=frozenset({"status", "priority", "pattern", "rationale"}))}

```
{core.get('pattern') or '（待填写）'}
```

**依据**：{core.get('rationale') or '（待填写）'}

---

## 5. 双缓冲策略 {_section_status_badge(dbuf)}

| 项 | 值 |
|----|-----|
{_kv_table(dbuf, skip=frozenset({"status", "priority", "rationale"}))}
**依据**：{dbuf.get('rationale') or '（待填写）'}

---

## 6. LDM 使用规划 {_section_status_badge(ldm)}

| 项 | 值 |
|----|-----|
| budget_bytes_per_spe | {ldm.get('budget_bytes_per_spe')} |
| total_bytes | {ldm.get('total_bytes', 0)} |
| fits_budget | {ldm.get('fits_budget', True)} |

{_render_buffer_table(ldm.get('buffers', []))}
**备注**：{ldm.get('allocation_notes') or '（待填写）'}

---

## 7. 流水排布 {_section_status_badge(pipe)}

**数据流**：{pipe.get('dataflow_description') or '（待填写）'}

### Pass 列表

{_bullets(pipe.get('passes', []))}

### 同步点

{_bullets(pipe.get('sync_points', []))}

---

## 8. teco-ops 落地与 Golden

```bash
python3 ops/kernel-bench/scripts/teco_ops_test.py \\
  --op {op} --demo-root ../../teco-ops --build-only --json
```

| 字段 | 值 |
|------|-----|
| scaffold.family | `{scaffold.get('family')}` |
| scaffold.smoke | {scaffold.get('smoke')} |
| scaffold.patch_runsh | {scaffold.get('patch_runsh')} |
| scaffold.upgrade_common | {scaffold.get('upgrade_common')} |
| golden.framework | `{golden.get('framework')}` |
| golden.atol | {golden.get('atol')} |
| golden.rtol | {golden.get('rtol')} |
{_render_extensions(doc)}"""
